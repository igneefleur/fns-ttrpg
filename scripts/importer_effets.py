"""Importe le référentiel et les consommables JDR du classeur, sans accès réseau.

Usage : python scripts/importer_effets.py /chemin/equipment.xlsx
Les lignes sans conversion restent descriptives : aucune règle n'est déduite
de leur texte Outward pendant l'exécution de la fiche.
"""
import json,re,sys,unicodedata
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
def slug(v):return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFD',str(v)).encode('ascii','ignore').decode().lower()).strip('-')
def duration(v):
 m=re.fullmatch(r'(\d+) (seconde|minute|heure)s?',str(v))
 return int(m[1])*{'seconde':1,'minute':60,'heure':3600}[m[2]] if m else 0 if v=='Instantané' else None
def build(path):
 w=openpyxl.load_workbook(path,read_only=False,data_only=True);effects=[]
 elements=['feu','froid','eclair','decomposition','ethere'];romans={'I':1,'II':2,'III':3,'IV':4,'V':5}
 for row in list(w['Effets'].values)[6:]:
  if not row or not row[0]:continue
  n,f,c,r,v,d,cu,note= list(row[:8]);op={};rank=romans.get(str(r),0);v=v or ''
  rec=re.search(r'(?:Niveau TTRPG|niveau) (\d+)',v)
  if rec:
   level=int(rec[1]);points=[0,5,10,15,20][(level-1)%4+1] if level else 0
   if not rank:rank=level
   cadence=28800 if level<=4 else 3600 if level<=8 else 600 if level<=12 else 60 if level<=16 else 3
   res='pi' if 'Innocence' in v else 'pe' if ' PE' in v or '(PE)' in v else 'pm' if ' PM' in v or '(PM)' in v else 'pv'
   op['regen']={'reserve':res,'niveau':level,'points':points,'cadence':cadence}
  if n.startswith('Effondrement ') and 'ignoré' in n:
   k='ps' if 'satiété' in n else 'ph' if 'hydratation' in n else 'pr' if 'sommeil' in n else 'total';op['effondrement']={k:-int(str(r))}
  boon={'Bénédiction':'eclair','Fraîcheur':'froid','Brume':'ethere','Possession':'decomposition','Chaleur':'feu'}
  if n in boon:
   op['degats']={boon[n]:20};op['resistances']={boon[n]:20}
   if n=='Fraîcheur':op['climat']={'chaud':8}
   if n=='Chaleur':op['climat']={'froid':8}
  if n in ['Rage','Impact augmenté']:op['paradeAdverse']=2
  if n=='Discipline':op['attaque']=-1
  if n.startswith('Attaque physique augmentée'):
   rank=10 if '+10 %' in n else 20;op['degats']={'physique':rank}
  if n=='Scintillement':op['degats']={k:15 for k in elements}
  if n=='Résistance aux impacts augmentée':op['parade']=-2
  if n.startswith('Protection augmentée'):op['protections']={'physique':30 if rank==2 else 10}
  if n.startswith('Barrière augmentée'):op['protections']={k:30 if rank==2 else 10 for k in elements}
  if n in ['Résistance élémentaire','Immunité élémentaire']:op['resistances' if n=='Résistance élémentaire' else 'immunites']={k:20 if n=='Résistance élémentaire' else 100 for k in elements}
  if n=='Vitesse augmentée':op['pas']=3
  if n=='Discrétion augmentée':op['sensAdverse']=2
  if n=='Réduction du coût en mana':op['couts']={'pm':-15}
  if n=='Protection climatique':op['climat']={'froid':20,'chaud':20};op['resistances']={'decomposition':10}
  if n.startswith('Résistance à la corruption'):op['innocence']=50 if rank==2 else 20
  if n.startswith('Résistance aux altérations '):op['alterations']=75 if 'augmentée' in n else -30
  if n=='Résistance de la lampe du Djinn':op['alterations']=15
  if n=='Protection de la lanterne des âmes':op['resistances']={k:20 if k=='ethere' else 10 for k in elements}
  if n.startswith('Alerte +'):op['alerte']=2 if '+2' in n else 1
  if n=='Effet du Pacificateur':op['permanent']={'pv':20,'pe':20,'pm':20}
  if n=='Énergisé':op['finTour']=1
  if n=='Protection contre le froid':op['climat']={'froid':15}
  if n=='Protection contre la chaleur':op['climat']={'chaud':15}
  if n=='Protection chaleur des eaux':op['climat']={'chaud':10}
  if n=='Effondrement d’alcool +1':op['alcool']=1
  suppressions={'Suppression du poison':['Poison'],'Suppression du saignement':['Saignement'],'Suppression de Burning':['Burning'],'Suppression d’Immolate':['Immolate'],'Suppression d’Affaibli':['Affaibli'],'Suppression de Sapped':['Sapped'],"Suppression d'Affaibli et Sapped":['Affaibli','Sapped'],'Dissipation des maléfices':['Chill','Curse','Doomed','Haunted','Scorched','Douleur','Confusion','The Hunch']}
  if n in suppressions:op['supprime']=suppressions[n]
  if n=='Panacée':op['panacee']=True
  if n.startswith('Imprégnation '):
   op['impregnation']={};target=next((k for k,word in [('feu','feu'),('froid','glace'),('eclair','foudre'),('decomposition','décomposition'),('ethere','éthérée')] if word in n),None)
   if target:
    pct=re.search(r'(\d+) %',v);flat=re.search(r'plus (\d+) dégâts',v)
    op['impregnation']={'type':target,'pct':int(pct[1]) if pct else 0,'flat':int(flat[1]) if flat else 0}
  auto=bool(op);manual=c in ['Maladie','Traitement','Altération','Référence Outward'] or not auto
  effects.append({'id':n,'nom':n,'famille':f,'categorie':c,'rang':rank,'description':v,'note':note or '', 'duree':duration(d),'dureeTexte':d,'cumul':cu,'ops':op,'manuel':manual,'negatif':c in ['Malus','Maladie','Altération'],'dissipable':c not in ['Permanent','Maladie','Référence Outward'],'etat':row[8] if len(row)>8 else 'Défini'})
 # Les états supprimés par les consommables doivent aussi pouvoir être ajoutés manuellement.
 known={x['id'] for x in effects}
 for n in ['Saignement','Saignement extrême','Burning','Immolate','Affaibli','Curse','Doomed','Haunted','Scorched','The Hunch']:
  if n not in known:effects.append({'id':n,'nom':n,'famille':'Saignement' if n.startswith('Saignement') else n,'categorie':'Altération','rang':0,'description':'Conséquences gérées par le joueur.','note':'','duree':None,'dureeTexte':'Durée manuelle','cumul':'Non','ops':{},'manuel':True,'negatif':True,'dissipable':True,'etat':'Manuel'})
 templates=[]
 for sheet,first,last,start,fx,rest in [('Nourriture',8,153,0,19,25),('Potions',7,50,0,8,23)]:
  s=w[sheet]
  for r in range(first,last+1):
   def cell(c):return s.cell(r,c).value
   name=cell(1)
   if not name:continue
   liquid=sheet=='Potions' or str(cell(3)).lower()=='boisson';paste=any(t in str(name).lower() for t in ['jam','paste','confiture'])
   item={'id':'tableur:'+slug(sheet+'-'+name),'nom':name,'qte':1,'typeObjet':'liquide' if liquid else 'pate' if paste else 'solide','sousType':'ingerable','nourri':True,'contenu':'liquide' if liquid else 'pate' if paste else '', 'effets':[cell(c) for c in range(fx,fx+5) if cell(c)],'restaurations':{},'unitesRestauration':{},'conditionsEffets':{},'ou':'poches','emp':-1,'dans':''}
   if sheet=='Nourriture':item.update(poids=cell(4) or 0,places=cell(7) or 0,tempsConsommation=cell(8) or 0,desc=cell(13) or '',img=cell(14) or '',volume=(cell(4) or 0)/(cell(11) or 1),ps=cell(5) or 0,ph=cell(6) or 0)
   else:item.update(poids=cell(3) or 0,volume=cell(4) or 0,places=cell(5) or 0,ph=cell(6) or 0,ps=0,tempsConsommation=cell(7) or 0,achat=cell(18) or 0,vente=cell(19),desc=cell(20) or '',img=cell(21) or '')
   for j,k in enumerate(['pv','pe','pm','pr','ps','ph','pi']):
    item['restaurations'][k]=cell(rest+j*2) or 0;item['unitesRestauration'][k]=cell(rest+j*2+1) or 'points'
   if sheet=='Nourriture':
    conditions=str(cell(39) or '')
    for part in conditions.split(' ; '):
     if ' : ' in part:
      k,val=part.split(' : ',1);m=re.search(r'≤ (\d+)',val)
      if m:item['conditionsEffets'][k]=int(m[1])
   item['encombre']=item['volume'];templates.append(item)
 gear=[];caracs={'FOR':'Force','DEX':'Dexterite','INT':'Intelligence','FER':'Ferveur','VIG':'Vigueur','END':'Endurance','RES':'Resistance','CHA':'Chance'}
 weapon_names=['Dagues','Gantelets','Rapières','Sabres','Épées 1 main','Haches 1 main','Masses 1 main','Bâtons','Lances','Hallebardes','Grands sabres','Grandes épées','Grandes haches','Faux de guerre','Grands marteaux','Fléaux à deux mains']
 keys=['dague','gantelet','rapiere','sabre','epee','hache','masse','baton','lance','hallebarde','grand-sabre','grande-epee','grande-hache','faux-de-guerre','grand-marteau','fleau']
 for sheet in ['Backpacks']+weapon_names+['Tête','Haut','Mains','Bas','Pieds','Robes']:
  s=w[sheet];header=15 if sheet=='Backpacks' else 12;headers=[c.value for c in s[header]]
  for row in list(s.values)[header:]:
   if not row or not row[0]:continue
   get=lambda name,default=None:row[headers.index(name)] if name in headers and headers.index(name)<len(row) else default
   item={'nom':row[0],'qte':1,'ou':'poches','emp':-1,'dans':'','effets':[],'taille':get('Taille',''),'volume':get('Volume (L)',0) or 0,'img':get('Lien de l’image',get('Lien image',get('Lien de l’image', ''))) or ''}
   if sheet=='Backpacks':
    item.update(typeObjet='rangement',sousType='sac',sac=True,poids=get('Poids (kg)',0),construction=get('Type',''),cap=get('Capacité (L)',0),tailleInterieure=get('Taille max. acceptée',''),tailleAccroche='Moyenne',ep=get('Accroches',0),achat=get('Prix d’achat (argent)',0),vente=get('Prix de vente (argent)'),desc=get('Bonus spécial',''))
   elif sheet in weapon_names:
    rep={k:float(get(h,0) or 0)*100 for k,h in [('physique','Physique (%)'),('feu','Feu (%)'),('froid','Froid (%)'),('eclair','Éclair (%)'),('decomposition','Décomposition (%)'),('ethere','Éthéré (%)')]}
    a={'attaque':str(get('Attaque',0)),'degats':str(get('Dégâts de base',0)),'parade':str(get('Parade',0)),'reduction':str(get('Réduction de base',0)),'comp':'','type':keys[weapon_names.index(sheet)],'mains':1 if weapon_names.index(sheet)<7 else 2,'repartition':rep,'hits':get('Hits',1),'dex':get('DEX requis',0),'for':get('FOR requis',0),'modsDegats':[],'modsParade':[]}
    for h,val in zip(headers,row):
     if h and str(h).startswith('Mod dégâts'):a['modsDegats'].append(caracs.get(str(val),str(val) if val in caracs.values() else ''))
     if h and str(h).startswith('Mod réduction'):a['modsParade'].append(caracs.get(str(val),str(val) if val in caracs.values() else ''))
    item.update(typeObjet='arme',sousType=a['type'],arme=a,poids=get('Poids (kg)',0),desc=get('Description',''),achat=get('Prix d’achat (argent)',0),vente=get('Prix de vente (argent)'))
   else:
    item.update(typeObjet='vetement',sousType={'Tête':'tete','Haut':'haut','Mains':'mains','Bas':'bas','Pieds':'pieds','Robes':'hautbas'}[sheet],vet={'Tête':'tete','Haut':'haut','Mains':'mains','Bas':'bas','Pieds':'pieds','Robes':'hautbas'}[sheet],construction=get('Construction',''),poids=get('Poids (kg)',0),achat=get('Prix achat',0),vente=get('Prix vente'),desc=get('Description',''),prot_physique=get('Protection',0),froid=get('Défense météo froid',0),chaud=get('Défense météo chaleur',0),stats={})
    for k,word in [('physique','physique'),('feu','feu'),('froid','froid'),('eclair','foudre'),('decomposition','corruption'),('ethere','éthéré')]:
     for h,val in zip(headers,row):
      if h and str(h).startswith('Rés.') and word in str(h).lower() and '%' in str(h) and 'personnage' not in str(h):item['res_'+k]=float(val or 0)
      if h and str(h).startswith('Bonus dégâts') and word in str(h).lower():item['stats'].setdefault('degats',{})[k]=float(val or 0)
    for res,label in [('pm','Coût mana (%)'),('pe','Coût endurance (%)')]:item['stats'].setdefault('couts',{})[res]=float(get(label,0) or 0)
    item['stats']['innocence']=float(get('Rés. corruption personnage (%)',0) or 0)
    item['stats']['alterations']=float(get('Rés. statuts (%)',0) or 0)
    item['stats']['vitesse']=float(get('Vitesse déplacement (%)',0) or 0)
    barrier=get('Barrière',0) or 0
    for k in elements:item['prot_'+k]=barrier
   item['encombre']=item['volume'];item['id']='tableur:'+slug(sheet+'-'+item['nom']);gear.append(item)
 for sheet in ['Lanternes et torches','Tentes et couchages']:
  for row in list(w[sheet].values)[7:]:
   if not row or not row[0]:continue
   if sheet=='Lanternes et torches' and str(row[1]).lower() not in ['torche','lanterne','torche-lanterne']:continue
   if len(row)<15:continue
   lantern=sheet=='Lanternes et torches';fx=row[20] if lantern and len(row)>20 else row[7] if not lantern else None
   item={'id':'tableur:'+slug(sheet+'-'+row[0]),'nom':row[0],'qte':1,'poids':row[2] or 0,'volume':row[3] or 0,'encombre':row[3] or 0,'taille':row[4] or '', 'construction':'Souple' if not lantern else '', 'typeObjet':'eclairage' if lantern else 'couchage','sousType':('torche' if 'Torch' in row[0] else 'lanterne') if lantern else ('couchage' if 'Bedroll' in row[0] else 'tente'),'effets':[fx] if fx else [],'achat':row[11] or 0,'vente':row[12],'desc':row[13] or '', 'img':row[14] or '', 'ou':'poches','dans':'','emp':-1,'proprietes':{'autonomie':row[18] if lantern and len(row)>18 else None,'accrochable':row[8] if lantern else None,'recharge':row[7] if lantern else None}}
   gear.append(item)
 for r in [8,9,10,11,12,13,16,17,18,19]:
  s=w['Imprégnations'];row=[s.cell(r,c).value for c in range(1,19)]
  if not row[0]:continue
  liquid='Varnish' in row[0];templates.append({'id':'tableur:'+slug(row[0]),'nom':row[0],'qte':1,'typeObjet':'liquide' if liquid else 'solide','sousType':'consommable','nourri':False,'poids':row[2] or 0,'volume':row[12] or 0,'encombre':row[12] or 0,'taille':row[13] or '', 'effets':[row[3]] if row[3] else [],'tempsConsommation':row[9] or 1,'achat':row[14] or 0,'vente':row[15],'desc':row[16] or '', 'img':row[17] or '', 'ou':'poches','dans':'','emp':-1})
 out={'version':1,'source':'ttrpg outward equipment','effets':effects,'objets':templates,'equipements':gear}
 dest=ROOT/'docs/data/owd-effets.json';dest.parent.mkdir(exist_ok=True);dest.write_text(json.dumps(out,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
 print(len(effects),'effets,',len(templates),'modèles de consommables importés')
if __name__=='__main__':build(sys.argv[1])
